"""
PolicySetu DOCX Document Parser.
Extracts structured paragraphs, headings, lists, tables, and document metadata from Word (.docx) files.
Preserves table structures explicitly as TableContent without destructive string flattening.
"""

import hashlib
import io
from pathlib import Path
from typing import Any, Dict, List, Optional
import docx  # type: ignore

try:
    from .models import (
        DocumentContent,
        DocumentExtractionMethod,
        DocumentProcessingStatus,
        PageContent,
        TableContent,
        TextBlock,
    )
    from .provenance import DocumentType
except (ImportError, ValueError):
    from src.documents.models import (
        DocumentContent,
        DocumentExtractionMethod,
        DocumentProcessingStatus,
        PageContent,
        TableContent,
        TextBlock,
    )
    from src.documents.provenance import DocumentType


class DocxParser:
    """
    Structured Word Document (.docx) parser.
    Preserves hierarchical headings, paragraphs, and multi-column tables.
    """

    def parse_bytes(
        self,
        docx_bytes: bytes,
        filename: str = "document.docx",
        document_id: Optional[str] = None,
    ) -> DocumentContent:
        """Parses in-memory DOCX bytes."""
        sha256 = hashlib.sha256(docx_bytes).hexdigest()
        doc_id = document_id or f"doc_{sha256[:12]}"

        try:
            doc = docx.Document(io.BytesIO(docx_bytes))
        except Exception as exc:
            return DocumentContent(
                document_id=doc_id,
                file_path=filename,
                file_name=filename,
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                document_type=DocumentType.UNKNOWN_DOCUMENT,
                status=DocumentProcessingStatus.CORRUPTED,
                sha256=sha256,
                error_message=f"Failed to parse DOCX archive: {exc}",
            )

        text_blocks: List[TextBlock] = []
        full_text_parts: List[str] = []

        # 1. Extract paragraphs and headings
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                style_name = str(p.style.name) if (p.style and p.style.name) else "Normal"
                prefix = f"[{style_name}] " if "Heading" in style_name else ""
                formatted_text = f"{prefix}{text}"
                text_blocks.append(
                    TextBlock(
                        text=formatted_text,
                        page_number=1,
                        confidence=1.0,
                        extraction_method=DocumentExtractionMethod.DOCX_XML,
                    )
                )
                full_text_parts.append(formatted_text)

        # 2. Extract structured tables
        tables: List[TableContent] = []
        for tbl in doc.tables:
            if not tbl.rows:
                continue
            headers = [cell.text.strip() for cell in tbl.rows[0].cells]
            rows_data = []
            for row in tbl.rows[1:]:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):
                    rows_data.append(row_cells)

            if headers or rows_data:
                tables.append(TableContent(headers=headers, rows=rows_data, page_number=1))

        # 3. Extract core properties metadata
        props = doc.core_properties
        metadata = {
            "title": props.title,
            "author": props.author,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
            "format": "DOCX",
        }

        # Pack into PageContent (DOCX does not have native digital page breaks without Word rendering)
        page = PageContent(
            page_number=1,
            text="\n".join(full_text_parts),
            text_blocks=text_blocks,
            tables=tables,
            is_scanned=False,
        )

        return DocumentContent(
            document_id=doc_id,
            file_path=filename,
            file_name=filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            document_type=DocumentType.UNKNOWN_DOCUMENT,
            pages=[page],
            tables=tables,
            metadata=metadata,
            status=DocumentProcessingStatus.VALID,
            sha256=sha256,
        )

    def parse_file(self, file_path: str) -> DocumentContent:
        """Parses a local DOCX file path."""
        p = Path(file_path)
        raw_bytes = p.read_bytes()
        sha256 = hashlib.sha256(raw_bytes).hexdigest()
        doc_id = f"doc_{sha256[:12]}"
        return self.parse_bytes(raw_bytes, filename=p.name, document_id=doc_id)
