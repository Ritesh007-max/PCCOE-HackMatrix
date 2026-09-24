"""
Unit tests for Document Parsers, ScanDetector, DocumentTypeDetector, and DocumentPipeline.
"""

import io
import unittest
from PIL import Image
import pymupdf  # type: ignore
import docx  # type: ignore

from src.documents.models import DocumentProcessingStatus, DocumentExtractionMethod
from src.documents.provenance import DocumentType
from src.documents.scan_detector import ScanDetector, PageType
from src.documents.detector import DocumentTypeDetector
from src.documents.pdf_parser import LayeredPDFParser
from src.documents.docx_parser import DocxParser
from src.documents.image_parser import ImageParser
from src.documents.pipeline import DocumentPipeline


class TestDocumentParsers(unittest.TestCase):

    def test_scan_detector_classification(self):
        detector = ScanDetector(min_char_threshold=30)

        # 1. Native text page (high char count, no image)
        res1 = detector.analyze_page(1, "This is a digital text page with ample administrative text content.", image_count=0)
        self.assertEqual(res1.page_type, PageType.NATIVE_TEXT)
        self.assertFalse(res1.is_scanned)

        # 2. Scanned page (low text, 1 image)
        res2 = detector.analyze_page(2, "Seal", image_count=1)
        self.assertEqual(res2.page_type, PageType.SCANNED)
        self.assertTrue(res2.is_scanned)

        # 3. Hybrid page (adequate text + photo)
        res3 = detector.analyze_page(3, "Applicant Name: Ramesh Kumar; State: Gujarat; Income: 200000", image_count=1)
        self.assertEqual(res3.page_type, PageType.HYBRID)
        self.assertFalse(res3.is_scanned)

    def test_document_type_detector_unknown_support(self):
        detector = DocumentTypeDetector()

        # Aadhaar detection
        aadhaar_text = "Government of India Unique Identification Authority of India UIDAI 1234 5678 9012"
        dt1, conf1 = detector.detect(aadhaar_text)
        self.assertEqual(dt1, DocumentType.AADHAAR)
        self.assertGreater(conf1, 0.6)

        # Income Certificate detection
        income_text = "Office of the Tahsildar Certificate of Annual Family Income Rs 1,50,000/-"
        dt2, conf2 = detector.detect(income_text)
        self.assertEqual(dt2, DocumentType.INCOME_CERTIFICATE)

        # UNKNOWN Document must NOT be forced into official category
        random_text = "Shopping receipt for groceries, biscuits, and milk from local bakery."
        dt3, conf3 = detector.detect(random_text, filename="random_receipt.txt")
        self.assertIn(dt3, (DocumentType.UNKNOWN_DOCUMENT, DocumentType.UNKNOWN))
        self.assertEqual(conf3, 0.0)

    def test_pdf_layered_parser(self):
        # Create a synthetic in-memory PDF with PyMuPDF
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)
        page.insert_text((50, 72), "Government of Gujarat Domicile Certificate", fontsize=14)
        page.insert_text((50, 100), "Resident of Ahmedabad continuous residence 12 years", fontsize=11)
        pdf_bytes = doc.tobytes()
        doc.close()

        parser = LayeredPDFParser()
        content = parser.parse_bytes(pdf_bytes, filename="domicile.pdf")

        self.assertEqual(content.status, DocumentProcessingStatus.VALID)
        self.assertEqual(len(content.pages), 1)
        self.assertIn("Domicile Certificate", content.get_full_text())
        self.assertFalse(content.pages[0].is_scanned)

    def test_docx_parser_tables_and_headings(self):
        # Create a synthetic in-memory DOCX
        doc = docx.Document()
        doc.add_heading("Income Certificate Summary", level=1)
        doc.add_paragraph("Applicant: Priya Patel, Age: 21, Gender: Female")

        # Add a table
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Source"
        table.rows[0].cells[1].text = "Annual Amount (INR)"
        table.rows[1].cells[0].text = "Agriculture"
        table.rows[1].cells[1].text = "250000"

        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        parser = DocxParser()
        content = parser.parse_bytes(docx_bytes, filename="income_cert.docx")

        self.assertEqual(content.status, DocumentProcessingStatus.VALID)
        self.assertIn("Income Certificate Summary", content.get_full_text())
        self.assertEqual(len(content.tables), 1)
        self.assertEqual(content.tables[0].headers, ["Source", "Annual Amount (INR)"])
        self.assertEqual(content.tables[0].rows, [["Agriculture", "250000"]])

    def test_image_parser_and_pipeline(self):
        # Create a synthetic in-memory PNG image
        img = Image.new("RGB", (200, 100), color="white")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        png_bytes = buf.getvalue()

        pipeline = DocumentPipeline()
        content = pipeline.process_bytes(png_bytes, filename="scan_sample.png")

        self.assertEqual(content.status, DocumentProcessingStatus.VALID)
        self.assertEqual(content.mime_type, "image/png")
        self.assertEqual(len(content.pages), 1)
        self.assertTrue(content.pages[0].is_scanned)


if __name__ == "__main__":
    unittest.main()
